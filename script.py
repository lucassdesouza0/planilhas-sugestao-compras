from __future__ import annotations

import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter, column_index_from_string


@dataclass(frozen=True)
class Config:
    # Arquivos
    compras_xlsx: str           # arquivo A
    estoque_xlsx: str           # arquivo B
    output_xlsx: str

    # Abas
    compras_sheet_name: str = "Sugestão de Compras"
    estoque_sheet_name: str 

    # Cabeçalhos na aba de compras (nomes no cabeçalho da linha 1)
    product_code_header: str = "Cód. produto"
    category_header: str = "CATEGORIA"

    # Inserir colunas custo/custo_medio na posição fixa (letra)
    cost_insert_at_col_letter: str = "D"
    cost_headers: Tuple[str, str] = ("CUSTO", "CUSTO_MEDIO")

    # Mapeamento da aba Estoque (letras conforme você passou)
    # (código em C, custo em L, custo médio em M)
    stock_code_col_letter: str = "C"
    stock_cost_col_letter: str = "L"
    stock_avg_col_letter: str = "M"

    # Colunas a manter (mais seguro que “excluir”)
    keep_headers: Optional[List[str]] = None
    
    # Mapeamento para renomear cabeçalhos no arquivo final
    # {nome_original: nome_novo}
    header_renames: Optional[Dict[str, str]] = None

    uncategorized_sheet_name: str = "Sem categoria"

    # Fórmula: XLOOKUP (Excel 365) ou VLOOKUP (mais compatível)
    use_xlookup: bool = True


def normalize_header(s: Any) -> str:
    return str(s).strip().upper()


def read_header_map(ws: Worksheet) -> Dict[str, int]:
    header_map: Dict[str, int] = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        if cell.value is None:
            continue
        header_map[normalize_header(cell.value)] = col_idx
    return header_map


def iter_rows_as_dicts(ws: Worksheet, header_map: Dict[str, int]) -> List[Dict[str, Any]]:
    headers = list(header_map.keys())
    rows: List[Dict[str, Any]] = []
    for r in range(2, ws.max_row + 1):
        row_dict: Dict[str, Any] = {}
        empty = True
        for h in headers:
            v = ws.cell(row=r, column=header_map[h]).value
            row_dict[h] = v
            if v not in (None, ""):
                empty = False
        if not empty:
            rows.append(row_dict)
    return rows


def safe_sheet_title(wb: Workbook, title: str) -> str:
    bad_chars = ['\\', '/', '*', '?', ':', '[', ']']
    t = str(title).strip()
    for ch in bad_chars:
        t = t.replace(ch, "-")
    t = t[:31] if len(t) > 31 else t
    if not t:
        t = "Categoria"

    base = t
    i = 2
    while t in wb.sheetnames:
        suffix = f" {i}"
        t = (base[:31 - len(suffix)] + suffix)[:31]
        i += 1
    return t


def build_lookup_formula(cfg: Config, stock_sheet_name_in_output: str, product_code_cell_ref: str, return_col_letter: str) -> str:
    if cfg.use_xlookup:
        # =XLOOKUP(A2, Estoque!$C:$C, Estoque!$L:$L, "")
        return (
            f'=XLOOKUP({product_code_cell_ref}, '
            f'\'{stock_sheet_name_in_output}\'!${cfg.stock_code_col_letter}:${cfg.stock_code_col_letter}, '
            f'\'{stock_sheet_name_in_output}\'!${return_col_letter}:${return_col_letter}, "")'
        )

    # VLOOKUP compatível: range C:M (C=1, L=10, M=11 dentro do range)
    start = cfg.stock_code_col_letter
    end = cfg.stock_avg_col_letter  # M
    start_i = column_index_from_string(start)
    return_i = column_index_from_string(return_col_letter)
    col_index_within_range = return_i - start_i + 1
    return (
        f'=IFERROR(VLOOKUP({product_code_cell_ref}, '
        f'\'{stock_sheet_name_in_output}\'!${start}:${end}, {col_index_within_range}, FALSE), "")'
    )


def write_sheet(ws: Worksheet, headers: List[str], rows: List[Dict[str, Any]], cfg: Config, stock_sheet_name_in_output: str):
    # Cabeçalho (aplicando renomeações se configuradas)
    display_headers = []
    for h in headers:
        if cfg.header_renames and normalize_header(h) in {normalize_header(k) for k in cfg.header_renames}:
            # Encontra o nome original correspondente (case-insensitive)
            for old_name, new_name in cfg.header_renames.items():
                if normalize_header(h) == normalize_header(old_name):
                    display_headers.append(new_name)
                    break
        else:
            display_headers.append(h)
    
    for col_idx, h in enumerate(display_headers, start=1):
        ws.cell(row=1, column=col_idx).value = h

    # Índices das colunas relevantes na aba resultante
    code_col_idx = headers.index(cfg.product_code_header) + 1

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, h in enumerate(headers, start=1):
            if h == cfg.cost_headers[0]:
                code_cell = f"{get_column_letter(code_col_idx)}{r_idx}"
                ws.cell(row=r_idx, column=c_idx).value = build_lookup_formula(
                    cfg, stock_sheet_name_in_output, code_cell, cfg.stock_cost_col_letter
                )
            elif h == cfg.cost_headers[1]:
                code_cell = f"{get_column_letter(code_col_idx)}{r_idx}"
                ws.cell(row=r_idx, column=c_idx).value = build_lookup_formula(
                    cfg, stock_sheet_name_in_output, code_cell, cfg.stock_avg_col_letter
                )
            else:
                ws.cell(row=r_idx, column=c_idx).value = row.get(normalize_header(h))


def copy_sheet_values(src_ws: Worksheet, dst_ws: Worksheet):
    # copia valores e fórmulas (estilos não são essenciais para lookup)
    for r in range(1, src_ws.max_row + 1):
        for c in range(1, src_ws.max_column + 1):
            dst_ws.cell(row=r, column=c).value = src_ws.cell(row=r, column=c).value


def _normalize_code(v: Any) -> str:
    return str(v).strip()


def _is_zeroish(v: Any) -> bool:
    eps = 1e-9
    if v is None:
        return True
    if isinstance(v, (int, float)):
        return abs(float(v)) < eps
    s = str(v).strip().replace(",", ".")
    if s == "":
        return True
    try:
        return abs(float(s)) < eps
    except ValueError:
        return False


def copy_sheet_values_dedup_cost_zero(src_ws: Worksheet, dst_ws: Worksheet, cfg: Config) -> None:
    """
    Copia a aba Estoque para o output, removendo duplicados por código.

    Regra: se houver registros duplicados para o mesmo código, e algum deles tiver
    custo (coluna cfg.stock_cost_col_letter) = 0, prioriza manter o registro com
    custo != 0. Caso todos sejam 0, mantém o primeiro encontrado.

    Isso evita o XLOOKUP retornar o registro "errado" por pegar a primeira ocorrência.
    """
    max_col = src_ws.max_column

    # Header (linha 1)
    for c in range(1, max_col + 1):
        dst_ws.cell(row=1, column=c).value = src_ws.cell(row=1, column=c).value

    code_idx = column_index_from_string(cfg.stock_code_col_letter)
    cost_idx = column_index_from_string(cfg.stock_cost_col_letter)

    # code -> (row_values, is_cost_zeroish)
    selected: Dict[str, Tuple[List[Any], bool]] = {}
    order: List[str] = []
    duplicates_with_zero_removed = 0
    duplicates_kept_first = 0

    for r in range(2, src_ws.max_row + 1):
        code_raw = src_ws.cell(row=r, column=code_idx).value
        if code_raw in (None, ""):
            continue
        code = _normalize_code(code_raw)
        if not code:
            continue

        row_values = [src_ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        cost_val = src_ws.cell(row=r, column=cost_idx).value
        cost_is_zero = _is_zeroish(cost_val)

        if code not in selected:
            selected[code] = (row_values, cost_is_zero)
            order.append(code)
            continue

        _, prev_cost_is_zero = selected[code]
        if prev_cost_is_zero and not cost_is_zero:
            # Substitui o registro "0" por um registro com custo != 0
            selected[code] = (row_values, cost_is_zero)
            duplicates_with_zero_removed += 1
        else:
            # Mantém o primeiro (se ambos 0, ou ambos != 0, ou já temos != 0 e o novo é 0)
            duplicates_kept_first += 1

    # Escreve linhas deduplicadas
    out_r = 2
    for code in order:
        row_values, _ = selected[code]
        for c, v in enumerate(row_values, start=1):
            dst_ws.cell(row=out_r, column=c).value = v
        out_r += 1

    if duplicates_with_zero_removed or duplicates_kept_first:
        print(
            f"Estoque cleanup: deduplicado por código; "
            f"substituições por custo!=0: {duplicates_with_zero_removed}; "
            f"duplicados ignorados: {duplicates_kept_first}"
        )


def main(cfg: Config) -> None:
    compras_path = Path(cfg.compras_xlsx).expanduser().resolve()
    estoque_path = Path(cfg.estoque_xlsx).expanduser().resolve()
    output_path = Path(cfg.output_xlsx).expanduser().resolve()

    if not compras_path.exists():
        raise FileNotFoundError(f"Arquivo de compras não encontrado: {compras_path}")
    if not estoque_path.exists():
        raise FileNotFoundError(f"Arquivo de estoque não encontrado: {estoque_path}")

    wb_compras = load_workbook(compras_path)
    wb_estoque = load_workbook(estoque_path)

    if cfg.compras_sheet_name not in wb_compras.sheetnames:
        raise ValueError(f"Aba '{cfg.compras_sheet_name}' não existe no arquivo de compras. Abas: {wb_compras.sheetnames}")
    if cfg.estoque_sheet_name not in wb_estoque.sheetnames:
        raise ValueError(f"Aba '{cfg.estoque_sheet_name}' não existe no arquivo de estoque. Abas: {wb_estoque.sheetnames}")

    ws_compras = wb_compras[cfg.compras_sheet_name]
    ws_estoque = wb_estoque[cfg.estoque_sheet_name]

    header_map = read_header_map(ws_compras)

    code_h_norm = normalize_header(cfg.product_code_header)
    cat_h_norm = normalize_header(cfg.category_header)
    if code_h_norm not in header_map:
        raise ValueError(f"Cabeçalho de código '{cfg.product_code_header}' não encontrado na aba compras.")
    if cat_h_norm not in header_map:
        raise ValueError(f"Cabeçalho de categoria '{cfg.category_header}' não encontrado na aba compras.")

    rows = iter_rows_as_dicts(ws_compras, header_map)

    # Headers originais
    original_headers: List[str] = [str(c.value).strip() for c in ws_compras[1] if c.value is not None]

    # Decide o que manter
    if cfg.keep_headers:
        keep_norm = {normalize_header(h) for h in cfg.keep_headers}
        kept_headers = [h for h in original_headers if normalize_header(h) in keep_norm]
    else:
        kept_headers = original_headers[:]

    # Garante que código e categoria existem
    def ensure(headers: List[str], h: str) -> None:
        if normalize_header(h) not in {normalize_header(x) for x in headers}:
            headers.append(h)

    ensure(kept_headers, cfg.product_code_header)
    ensure(kept_headers, cfg.category_header)

    # Remove se já existirem custo/custo_medio
    kept_headers = [h for h in kept_headers if normalize_header(h) not in {normalize_header(cfg.cost_headers[0]), normalize_header(cfg.cost_headers[1])}]

    # Insere custo/custo_medio na posição fixada
    insert_at = min(column_index_from_string(cfg.cost_insert_at_col_letter), len(kept_headers) + 1)
    kept_headers.insert(insert_at - 1, cfg.cost_headers[0])
    kept_headers.insert(insert_at, cfg.cost_headers[1])

    # Agrupa por categoria
    by_category: Dict[str, List[Dict[str, Any]]] = {}
    for row in rows:
        cat_val = row.get(cat_h_norm)
        cat_name = str(cat_val).strip() if cat_val not in (None, "") else cfg.uncategorized_sheet_name
        by_category.setdefault(cat_name, []).append(row)

    # Workbook de saída
    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    # Cria abas por categoria
    for cat, cat_rows in sorted(by_category.items(), key=lambda x: x[0].lower()):
        ws_cat = wb_out.create_sheet(title=safe_sheet_title(wb_out, cat))
        write_sheet(ws_cat, kept_headers, cat_rows, cfg, stock_sheet_name_in_output=cfg.estoque_sheet_name)

    # Adiciona estoque como última aba (copiando de arquivo separado)
    stock_title = safe_sheet_title(wb_out, cfg.estoque_sheet_name)
    # Se o safe_sheet_title ajustou o nome (duplicado), precisamos usar o nome real no lookup.
    # Para simplificar: força o nome exatamente como cfg.estoque_sheet_name se estiver livre.
    if cfg.estoque_sheet_name not in wb_out.sheetnames:
        stock_title = cfg.estoque_sheet_name

    ws_out_stock = wb_out.create_sheet(title=stock_title)
    copy_sheet_values_dedup_cost_zero(ws_estoque, ws_out_stock, cfg)

    # IMPORTANTE: se stock_title != cfg.estoque_sheet_name, as fórmulas apontariam para nome diferente.
    # Por isso acima tentamos garantir o nome exato. Evite criar categoria chamada "Estoque".
    if stock_title != cfg.estoque_sheet_name:
        print(f"Atenção: aba estoque foi criada como '{stock_title}'. Evite categoria com nome 'Estoque'.")

    wb_out.save(output_path)
    print(f"OK: gerado {output_path}")


if __name__ == "__main__":
    # Uso:
    # python script.py compras.xlsx estoque.xlsx saida.xlsx
    if len(sys.argv) < 4:
        print("Uso: python script.py <compras.xlsx> <estoque.xlsx> <saida.xlsx>")
        sys.exit(1)

    cfg = Config(
        compras_xlsx=sys.argv[1],
        estoque_xlsx=sys.argv[2],
        output_xlsx=sys.argv[3],
        compras_sheet_name="Sugestão de Compras",
        estoque_sheet_name="Estoque 2026-01-27.xlsx",
        product_code_header="Cód. produto",
        category_header="CATEGORIA",
        cost_insert_at_col_letter="C",
        keep_headers=["Cód. produto", "Descrição","Saídas/Vendas","Qtd. Pontos de Venda","Estoque","Tempo de estoque", "Sugestão(Un.)", "Qtd. mínima(Un.)"],  # opcional
        header_renames={
            "Cód. produto": "Código",
            "Saídas/Vendas": "Saídas",
            "Qtd. Pontos de Venda": "Qtd. PDV",
            "Tempo de estoque": "Estoque(dias)"
        },
        use_xlookup=True,
    )
    main(cfg)
